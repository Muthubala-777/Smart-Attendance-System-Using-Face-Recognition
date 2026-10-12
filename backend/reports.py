import csv
import io
from datetime import datetime
from typing import List, Dict, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from backend.config import REPORTS_DIR


class AttendanceReportGenerator:
    """
    Multi-Format Attendance Report Generator.
    Produces high-fidelity:
    1. CSV
    2. Excel (.xlsx) with styled summary headers and status highlights
    3. PDF documents with institutional layout, metadata, summary statistics, and student registers
    """

    @staticmethod
    def generate_csv(records: List[Dict[str, Any]], title: str = "Attendance Report") -> io.StringIO:
        output = io.StringIO()
        writer = csv.writer(output)

        # Header Metadata
        writer.writerow(["SMART ATTENDANCE SYSTEM - ACADEMIC ATTENDANCE REPORT"])
        writer.writerow(["Title", title])
        writer.writerow(["Generated At", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        writer.writerow([])

        # Table Columns
        headers = [
            "S.No",
            "Student ID",
            "Student Name",
            "Department",
            "Year",
            "Section",
            "Subject",
            "Date",
            "Session Time",
            "Status",
            "First Seen",
            "Last Seen",
            "Similarity",
            "Liveness"
        ]
        writer.writerow(headers)

        for idx, r in enumerate(records, 1):
            writer.writerow([
                idx,
                r.get("student_id", ""),
                r.get("student_name", ""),
                r.get("department", ""),
                r.get("year", ""),
                r.get("section", ""),
                r.get("subject_name", ""),
                r.get("session_date", ""),
                f"{r.get('start_time', '')} - {r.get('end_time', '')}",
                r.get("status", "ABSENT"),
                r.get("first_seen", "--"),
                r.get("last_seen", "--"),
                f"{r.get('similarity_score', 0.0):.2f}" if r.get("similarity_score") else "--",
                r.get("liveness_status", "--")
            ])

        output.seek(0)
        return output

    @staticmethod
    def generate_excel(records: List[Dict[str, Any]], title: str = "Attendance Report") -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance Register"

        # Fonts & Styles
        title_font = Font(name="Calibri", size=15, bold=True, color="1E293B")
        meta_font = Font(name="Calibri", size=10, italic=True, color="64748B")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        
        present_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        present_font = Font(name="Calibri", size=10, bold=True, color="166534")
        absent_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        absent_font = Font(name="Calibri", size=10, bold=True, color="991B1B")

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # Title block
        ws.merge_cells("A1:N1")
        ws["A1"] = f"SMART ATTENDANCE SYSTEM — {title.upper()}"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:N2")
        ws["A2"] = f"Generated: {datetime.now().strftime('%d-%b-%Y %I:%M %p')} | Total Students: {len(records)}"
        ws["A2"].font = meta_font
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        # Column Headers
        headers = [
            "S.No", "Student ID", "Student Name", "Department", "Year", "Section",
            "Subject", "Date", "Session Time", "Status", "First Seen", "Last Seen",
            "Similarity", "Liveness"
        ]
        
        row_num = 4
        for col_num, h in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_num, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Data Rows
        present_count = 0
        for idx, r in enumerate(records, 1):
            row_num += 1
            status = r.get("status", "ABSENT")
            if status == "PRESENT":
                present_count += 1

            row_data = [
                idx,
                r.get("student_id", ""),
                r.get("student_name", ""),
                r.get("department", ""),
                r.get("year", ""),
                r.get("section", ""),
                r.get("subject_name", ""),
                r.get("session_date", ""),
                f"{r.get('start_time', '')} - {r.get('end_time', '')}",
                status,
                r.get("first_seen", "--"),
                r.get("last_seen", "--"),
                f"{r.get('similarity_score', 0.0):.2f}" if r.get("similarity_score") else "--",
                r.get("liveness_status", "--")
            ]

            for col_num, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")

                if col_num == 10:  # Status column
                    if status == "PRESENT":
                        cell.fill = present_fill
                        cell.font = present_font
                    else:
                        cell.fill = absent_fill
                        cell.font = absent_font

        from openpyxl.utils import get_column_letter
        # Auto-adjust column widths
        for col_idx, col in enumerate(ws.columns, 1):
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def generate_pdf(records: List[Dict[str, Any]], title: str = "Attendance Report") -> io.BytesIO:
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(letter),
            rightMargin=25,
            leftMargin=25,
            topMargin=25,
            bottomMargin=25
        )
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Heading1"],
            fontSize=16,
            textColor=colors.HexColor("#1E3A8A"),
            alignment=1,
            spaceAfter=4
        )
        sub_style = ParagraphStyle(
            "ReportSub",
            parent=styles["Normal"],
            fontSize=9,
            textColor=colors.HexColor("#64748B"),
            alignment=1,
            spaceAfter=12
        )

        elements = []
        elements.append(Paragraph("SMART ATTENDANCE SYSTEM — ACADEMIC REPORT", title_style))
        
        total_stu = len(records)
        present_count = sum(1 for r in records if r.get("status") == "PRESENT")
        pct = (present_count / max(1, total_stu)) * 100.0
        
        sub_text = f"Filter: {title} | Date: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Total: {total_stu} | Present: {present_count} | Absent: {total_stu - present_count} | Attendance: {pct:.1f}%"
        elements.append(Paragraph(sub_text, sub_style))

        # Table data
        headers = [
            "S.No", "ID", "Student Name", "Dept", "Yr/Sec", "Subject",
            "Date", "Slot", "Status", "First Seen", "Last Seen", "Conf", "Live"
        ]
        
        table_data = [headers]
        for idx, r in enumerate(records, 1):
            table_data.append([
                str(idx),
                str(r.get("student_id", "")),
                str(r.get("student_name", ""))[:20],
                str(r.get("department", "")),
                f"{r.get('year', '')}-{r.get('section', '')}",
                str(r.get("subject_name", ""))[:16],
                str(r.get("session_date", "")),
                f"{r.get('start_time', '')}-{r.get('end_time', '')}",
                str(r.get("status", "ABSENT")),
                str(r.get("first_seen", "--"))[-8:],
                str(r.get("last_seen", "--"))[-8:],
                f"{r.get('similarity_score', 0.0):.2f}" if r.get("similarity_score") else "--",
                str(r.get("liveness_status", "--"))[:5]
            ])

        t = Table(table_data, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('FONTSIZE', (0, 1), (-1, -1), 7.5),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 5),
            ('TOPPADDING', (0, 0), (-1, 0), 5),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white])
        ]))

        elements.append(t)
        doc.build(elements)
        output.seek(0)
        return output
